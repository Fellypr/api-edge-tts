import io
import os
import re
import json
import time
import uuid
import shutil
import asyncio
import logging
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, UploadFile, File, Form, HTTPException, BackgroundTasks, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from pypdf import PdfReader
from docx import Document
from mutagen.mp3 import MP3
import edge_tts

# ---------------------------------------------------------
# Configuração de Logs
# ---------------------------------------------------------
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s"
)
logger = logging.getLogger("gerador-audio")

# ---------------------------------------------------------
# Diretório de Armazenamento Temporário
# ---------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent
STORAGE_DIR = BASE_DIR / "storage"
STORAGE_DIR.mkdir(parents=True, exist_ok=True)

# Tempo limite para retenção de arquivos temporários (2 horas)
TEMPO_RETENCAO_SEGUNDOS = 2 * 3600

# ---------------------------------------------------------
# Inicialização do FastAPI
# ---------------------------------------------------------
app = FastAPI(
    title="API de Geração de Audiolivro e Sincronização de Texto",
    description="API assíncrona para extrair texto de PDF/DOCX, sintetizar áudio via Edge TTS e gerar mapa de tempo sincronizado por palavra com túnel SSE de logs.",
    version="2.0.0"
)

# Configuração de CORS para permitir requisições de qualquer front-end
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ---------------------------------------------------------
# Gerenciador de Túnel de Logs e Eventos em Tempo Real (SSE)
# ---------------------------------------------------------
class JobLogManager:
    """
    Gerencia canais de streaming SSE (Server-Sent Events) para múltiplos jobs.
    Permite aos clientes do frontend acompanhar logs e percentuais de progresso ao vivo.
    """
    def __init__(self):
        self.active_jobs: dict[str, asyncio.Queue] = {}
        self.job_history: dict[str, list[dict]] = {}
        self.job_status: dict[str, str] = {}

    def register_job(self, job_id: str) -> asyncio.Queue:
        queue = asyncio.Queue()
        self.active_jobs[job_id] = queue
        self.job_history[job_id] = []
        self.job_status[job_id] = "processando"
        return queue

    async def broadcast(self, job_id: str, tipo: str, dados: dict):
        evento = {
            "tipo": tipo,
            "job_id": job_id,
            "tempo": time.strftime("%H:%M:%S"),
            **dados
        }
        if job_id in self.job_history:
            self.job_history[job_id].append(evento)
        if job_id in self.active_jobs:
            await self.active_jobs[job_id].put(evento)

    async def log(self, job_id: str, mensagem: str, nivel: str = "INFO"):
        logger.info(f"[{job_id[:8]}] {mensagem}")
        await self.broadcast(job_id, "log", {"nivel": nivel, "mensagem": mensagem})

    async def progresso(self, job_id: str, porcentagem: int, parte_atual: int, total_partes: int, mensagem: str = ""):
        await self.broadcast(job_id, "progresso", {
            "porcentagem": porcentagem,
            "parte_atual": parte_atual,
            "total_partes": total_partes,
            "mensagem": mensagem
        })

    async def concluido(self, job_id: str, resultado: dict):
        self.job_status[job_id] = "concluido"
        await self.broadcast(job_id, "concluido", resultado)

    async def erro(self, job_id: str, mensagem: str):
        self.job_status[job_id] = "erro"
        logger.error(f"[{job_id[:8]}] {mensagem}")
        await self.broadcast(job_id, "erro", {"mensagem": mensagem})

    async def stream_events(self, job_id: str):
        """
        Gerador assíncrono para o endpoint SSE.
        Reproduz histórico inicial e transmite novos eventos em tempo real.
        """
        job_dir = STORAGE_DIR / job_id
        arquivo_info = job_dir / "info.json"

        # Se o job já concluiu anteriormente e não está mais ativo na memória
        if job_id not in self.active_jobs and arquivo_info.exists():
            with open(arquivo_info, "r", encoding="utf-8") as f:
                info = json.load(f)
            yield f"data: {json.dumps({'tipo': 'concluido', 'job_id': job_id, 'tempo': time.strftime('%H:%M:%S'), **info}, ensure_ascii=False)}\n\n"
            return

        if job_id not in self.active_jobs and not arquivo_info.exists():
            yield f"data: {json.dumps({'tipo': 'erro', 'job_id': job_id, 'mensagem': 'Job não encontrado ou expirado.'}, ensure_ascii=False)}\n\n"
            return

        # 1. Envia histórico de eventos já gerados antes do cliente conectar
        historico = list(self.job_history.get(job_id, []))
        for evt in historico:
            yield f"data: {json.dumps(evt, ensure_ascii=False)}\n\n"
            if evt.get("tipo") in ("concluido", "erro"):
                return

        queue = self.active_jobs.get(job_id)
        if not queue:
            return

        # 2. Escuta novos eventos em tempo real
        try:
            while True:
                try:
                    evento = await asyncio.wait_for(queue.get(), timeout=15.0)
                    yield f"data: {json.dumps(evento, ensure_ascii=False)}\n\n"
                    if evento.get("tipo") in ("concluido", "erro"):
                        break
                except asyncio.TimeoutError:
                    # Keep-alive ping para evitar timeout de proxies HTTP / Hugging Face Spaces
                    yield ": keep-alive\n\n"
        finally:
            # Mantém fila e histórico temporariamente para eventuais reconexões
            pass

    def cleanup(self, job_id: str):
        self.active_jobs.pop(job_id, None)
        self.job_history.pop(job_id, None)
        self.job_status.pop(job_id, None)

log_manager = JobLogManager()

# ---------------------------------------------------------
# Funções Auxiliares de Limpeza de Armazenamento
# ---------------------------------------------------------
def limpar_arquivos_antigos():
    """Remove pastas de jobs antigas para evitar estouro de disco."""
    agora = time.time()
    try:
        for job_folder in STORAGE_DIR.iterdir():
            if job_folder.is_dir():
                idade = agora - job_folder.stat().st_mtime
                if idade > TEMPO_RETENCAO_SEGUNDOS:
                    shutil.rmtree(job_folder, ignore_errors=True)
                    log_manager.cleanup(job_folder.name)
                    logger.info(f"Pasta temporária expirada removida: {job_folder.name}")
    except Exception as e:
        logger.error(f"Erro ao limpar arquivos antigos: {e}")

# ---------------------------------------------------------
# Extração e Limpeza de Textos
# ---------------------------------------------------------
def extrair_texto_pdf(conteudo_bytes: bytes, termo_cabecalho: Optional[str] = None) -> list[str]:
    """
    Lê o PDF e aplica filtros inteligentes:
    - Remove números de página soltos
    - Remove cabeçalhos repetitivos
    - Agrupa em frases e parágrafos fluidos
    """
    paragrafos_originais = []
    leitor = PdfReader(io.BytesIO(conteudo_bytes))
    termo_filtro = termo_cabecalho.strip().lower() if termo_cabecalho else None

    for pagina in leitor.pages:
        t = pagina.extract_text()
        if not t:
            continue

        linhas = t.split("\n")
        linhas_filtradas = []

        for linha in linhas:
            linha_limpa = linha.strip()

            # FILTRO 1: Ignora números de página soltos (ex: "12", "- 15 -", "iv")
            if re.match(r'^(?:\d{1,4}|[ivxlcdm]{1,6}|-\s*\d+\s*-)$', linha_limpa, re.IGNORECASE):
                continue

            # FILTRO 2: Ignora cabeçalhos repetitivos configurados
            if termo_filtro and termo_filtro in linha_limpa.lower() and len(linha_limpa) < 60:
                continue

            # FILTRO 3: Ignora linhas vazias ou ruídos muito curtos
            if len(linha_limpa) < 2:
                continue

            linhas_filtradas.append(linha_limpa)

        texto_pagina_limpo = " ".join(linhas_filtradas)

        # Agrupa em frases/parágrafos fluidos
        paragrafos = texto_pagina_limpo.split(". ")
        par_atual = ""
        for p in paragrafos:
            p_limpo = p.strip()
            if not p_limpo:
                continue
            par_atual += p_limpo + ". "
            if p_limpo.endswith((".", "!", "?", "”", '"', ":")):
                paragrafos_originais.append(par_atual.strip())
                par_atual = ""
        if par_atual.strip():
            paragrafos_originais.append(par_atual.strip())

    return paragrafos_originais

def extrair_texto_docx(conteudo_bytes: bytes) -> list[str]:
    """
    Lê o documento Word (.docx) e extrai parágrafos e células de tabelas.
    """
    paragrafos_originais = []
    doc = Document(io.BytesIO(conteudo_bytes))

    for p in doc.paragraphs:
        texto = p.text.strip()
        if len(texto) > 0:
            paragrafos_originais.append(re.sub(r'\s+', ' ', texto).strip())

    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                tc = re.sub(r'\s+', ' ', celula.text).strip()
                if len(tc) > 20 and tc not in paragrafos_originais:
                    paragrafos_originais.append(tc)

    return paragrafos_originais

def agrupar_em_lotes(paragrafos: list[str], tamanho_lote: int = 10000) -> list[str]:
    """
    Agrupa os parágrafos em partes de até tamanho_lote caracteres,
    preservando a quebra de parágrafo dupla ("\n\n").
    """
    partes = []
    bloco_atual = ""

    for par in paragrafos:
        if len(bloco_atual) + len(par) < tamanho_lote:
            bloco_atual += par + "\n\n"
        else:
            if bloco_atual.strip():
                partes.append(bloco_atual.strip())
            bloco_atual = par + "\n\n"

    if bloco_atual.strip():
        partes.append(bloco_atual.strip())

    return partes

# ---------------------------------------------------------
# Alinhamento de Palavras e Sincronização de Tempos
# ---------------------------------------------------------
def alinhar_pontuacao_e_paragrafos(texto_com_formatacao: str, tempos_brutos: list[dict]) -> list[dict]:
    """
    Alinha os tokens do texto original (com pontuação e marcadores de parágrafo)
    com os offsets retornados pelo Edge-TTS (WordBoundary).
    """
    if not tempos_brutos:
        return []

    tokens_originais = []
    paragrafos = texto_com_formatacao.split("\n\n")
    total_pars = len(paragrafos)

    for indice_par, paragrafo in enumerate(paragrafos):
        palavras_do_par = [p for p in paragrafo.split(" ") if p.strip()]
        total_pals = len(palavras_do_par)
        for idx_p, pal in enumerate(palavras_do_par):
            eh_fim_de_paragrafo = (idx_p == total_pals - 1) and (indice_par < total_pars - 1)
            texto_token = pal + ("\n\n" if eh_fim_de_paragrafo else "")
            limpa = re.sub(r'[^\w]', '', pal).lower()
            tokens_originais.append({"original": texto_token, "limpa": limpa})

    mapa_alinhado = []
    cursor_tempo = 0
    total_tempos = len(tempos_brutos)

    for token in tokens_originais:
        encontrou = False
        limite_busca = min(cursor_tempo + 5, total_tempos)
        for busca in range(cursor_tempo, limite_busca):
            palavra_tts_limpa = re.sub(r'[^\w]', '', tempos_brutos[busca]["palavra"]).lower()

            if token["limpa"] and (
                token["limpa"] == palavra_tts_limpa 
                or palavra_tts_limpa in token["limpa"] 
                or token["limpa"] in palavra_tts_limpa
            ):
                mapa_alinhado.append({
                    "palavra": token["original"],
                    "inicio": tempos_brutos[busca]["inicio"],
                    "fim": tempos_brutos[busca]["fim"]
                })
                cursor_tempo = busca + 1
                encontrou = True
                break

        if not encontrou:
            if mapa_alinhado:
                mapa_alinhado[-1]["palavra"] += " " + token["original"]
            elif total_tempos > 0:
                mapa_alinhado.append({
                    "palavra": token["original"],
                    "inicio": tempos_brutos[0]["inicio"],
                    "fim": tempos_brutos[0]["fim"]
                })

    return mapa_alinhado

# ---------------------------------------------------------
# Síntese Assíncrona via Edge-TTS
# ---------------------------------------------------------
async def gerar_parte_sincronizada(
    indice: int,
    total_partes: int,
    texto_parte: str,
    voz: str = "pt-BR-AntonioNeural",
    rate: str = "+0%",
    pitch: str = "-5Hz",
    job_id: Optional[str] = None
) -> tuple[bytes, list[dict], float]:
    """
    Sintetiza uma parte do texto com retry automático e captura eventos de WordBoundary.
    """
    for tentativa in range(1, 4):
        try:
            comunicador = edge_tts.Communicate(
                text=texto_parte,
                voice=voz,
                rate=rate,
                pitch=pitch,
                boundary="WordBoundary"
            )
            audio_bytes = bytearray()
            tempos_brutos = []

            async for chunk in comunicador.stream():
                if chunk["type"] == "audio":
                    audio_bytes.extend(chunk["data"])
                elif chunk["type"] == "WordBoundary":
                    inicio = chunk["offset"] / 10_000_000
                    duracao = chunk["duration"] / 10_000_000
                    tempos_brutos.append({
                        "palavra": chunk["text"],
                        "inicio": inicio,
                        "fim": inicio + duracao
                    })

            if not audio_bytes:
                raise ValueError("Stream de áudio retornou vazio.")

            mp3_info = MP3(io.BytesIO(audio_bytes))
            duracao_real_bloco = mp3_info.info.length

            tempos_formatados = alinhar_pontuacao_e_paragrafos(texto_parte, tempos_brutos)

            msg = f"Parte {indice + 1}/{total_partes} concluída ({duracao_real_bloco:.1f}s | {len(tempos_formatados)} palavras)"
            logger.info(msg)
            if job_id:
                await log_manager.log(job_id, msg)

            return bytes(audio_bytes), tempos_formatados, duracao_real_bloco

        except Exception as e:
            msg_aviso = f"Aviso: Tentativa {tentativa}/3 falhou para parte {indice + 1}: {e}"
            logger.warning(msg_aviso)
            if job_id:
                await log_manager.log(job_id, msg_aviso, nivel="WARNING")
            if tentativa < 3:
                await asyncio.sleep(5)

    msg_falha = f"Falha definitiva ao sintetizar parte {indice + 1}"
    logger.error(msg_falha)
    if job_id:
        await log_manager.log(job_id, msg_falha, nivel="ERROR")
    return b"", [], 0.0

# ---------------------------------------------------------
# Tarefa de Execução do Processamento (Segundo Plano)
# ---------------------------------------------------------
async def executar_processamento_livro(
    job_id: str,
    conteudo_bytes: bytes,
    extensao: str,
    nome_original: str,
    termo_cabecalho: Optional[str],
    voz: str,
    rate: str,
    pitch: str,
    tamanho_lote: int,
    retornar_mapa_inline: bool,
    base_url: str
) -> dict:
    job_dir = STORAGE_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)

    nome_base_livro = re.sub(r'[^\w\s-]', '', nome_original).strip().replace(' ', '_')
    if not nome_base_livro:
        nome_base_livro = "audiolivro"

    arquivo_mp3 = job_dir / f"{nome_base_livro}_Audiolivro.mp3"
    arquivo_mapa = job_dir / f"{nome_base_livro}_MapaDoTempo.json"
    arquivo_info = job_dir / "info.json"

    download_audio_url = f"{base_url}/download/{job_id}/audio"
    download_mapa_url = f"{base_url}/download/{job_id}/mapa"

    try:
        await log_manager.log(job_id, f"Iniciando processamento do livro '{nome_original}'")
        await log_manager.progresso(job_id, porcentagem=5, parte_atual=0, total_partes=0, mensagem="Extraindo texto do documento...")

        # 1. Extração do texto
        if extensao == ".pdf":
            paragrafos_originais = extrair_texto_pdf(conteudo_bytes, termo_cabecalho)
        else:
            paragrafos_originais = extrair_texto_docx(conteudo_bytes)

        if not paragrafos_originais:
            raise ValueError("Nenhum texto legível foi encontrado no arquivo fornecido.")

        await log_manager.log(job_id, f"Extração concluída: {len(paragrafos_originais)} parágrafos encontrados.")
        await log_manager.progresso(job_id, porcentagem=15, parte_atual=0, total_partes=0, mensagem="Dividindo texto em partes otimizadas...")

        # 2. Divisão em lotes
        partes = agrupar_em_lotes(paragrafos_originais, tamanho_lote=tamanho_lote)
        total_partes = len(partes)
        await log_manager.log(job_id, f"Texto estruturado e dividido em {total_partes} partes para síntese neural.")

        # 3. Processamento e síntese assíncrona
        mapa_completo = []
        tempo_acumulado = 0.0
        tamanho_grupo = 3

        with open(arquivo_mp3, "wb") as f_mp3:
            for i in range(0, total_partes, tamanho_grupo):
                fim_grupo = min(i + tamanho_grupo, total_partes)
                await log_manager.log(job_id, f"Processando lote de partes {i + 1} até {fim_grupo} de {total_partes}...")

                tarefas = [
                    gerar_parte_sincronizada(
                        indice=i + j,
                        total_partes=total_partes,
                        texto_parte=texto,
                        voz=voz,
                        rate=rate,
                        pitch=pitch,
                        job_id=job_id
                    )
                    for j, texto in enumerate(partes[i : fim_grupo])
                ]

                resultados = await asyncio.gather(*tarefas)

                for audio_bytes, tempos_parte, duracao_bloco in resultados:
                    if not audio_bytes:
                        continue

                    f_mp3.write(audio_bytes)

                    for item in tempos_parte:
                        mapa_completo.append({
                            "p": item["palavra"],
                            "i": round(item["inicio"] + tempo_acumulado, 3),
                            "f": round(item["fim"] + tempo_acumulado, 3)
                        })

                    tempo_acumulado += duracao_bloco

                porcentagem_calculada = int(20 + (fim_grupo / total_partes) * 70)
                await log_manager.progresso(
                    job_id=job_id,
                    porcentagem=porcentagem_calculada,
                    parte_atual=fim_grupo,
                    total_partes=total_partes,
                    mensagem=f"Sintetizadas {fim_grupo} de {total_partes} partes ({tempo_acumulado:.1f}s acumulados)"
                )

                if fim_grupo < total_partes:
                    await log_manager.log(job_id, "Aguardando 5 segundos para respeitar limites do serviço Edge-TTS...")
                    await asyncio.sleep(5)

        # Salva o mapa sincronizado em JSON
        await log_manager.log(job_id, "Salvando mapa temporal de sincronização e consolidando arquivos...")
        await log_manager.progresso(job_id, porcentagem=95, parte_atual=total_partes, total_partes=total_partes, mensagem="Consolidando arquivos...")

        with open(arquivo_mapa, "w", encoding="utf-8") as f_json:
            json.dump(mapa_completo, f_json, ensure_ascii=False, indent=2)

        info_job = {
            "status": "sucesso",
            "job_id": job_id,
            "nome_livro": nome_base_livro,
            "arquivo_original": nome_original,
            "duracao_segundos": round(tempo_acumulado, 2),
            "duracao_minutos": round(tempo_acumulado / 60, 2),
            "total_paragrafos": len(paragrafos_originais),
            "total_partes": total_partes,
            "total_palavras": len(mapa_completo),
            "criado_em": time.time(),
            "download_audio_url": download_audio_url,
            "download_mapa_url": download_mapa_url,
            "mp3_nome": arquivo_mp3.name,
            "json_nome": arquivo_mapa.name
        }

        with open(arquivo_info, "w", encoding="utf-8") as f_info:
            json.dump(info_job, f_info, ensure_ascii=False, indent=2)

        resposta = {
            "status": "sucesso",
            "job_id": job_id,
            "nome_livro": nome_base_livro,
            "duracao_segundos": round(tempo_acumulado, 2),
            "duracao_minutos": round(tempo_acumulado / 60, 2),
            "total_paragrafos": len(paragrafos_originais),
            "total_partes": total_partes,
            "total_palavras": len(mapa_completo),
            "download_audio_url": download_audio_url,
            "download_mapa_url": download_mapa_url
        }

        if retornar_mapa_inline:
            resposta["mapa"] = mapa_completo

        await log_manager.log(job_id, f"Processamento concluído com sucesso! Duração total: {round(tempo_acumulado / 60, 2)} min.")
        await log_manager.progresso(job_id, porcentagem=100, parte_atual=total_partes, total_partes=total_partes, mensagem="Audiolivro e mapa gerados com sucesso!")
        await log_manager.concluido(job_id, resposta)
        return resposta

    except Exception as e:
        logger.error(f"Erro durante o processamento do job {job_id}: {e}")
        shutil.rmtree(job_dir, ignore_errors=True)
        await log_manager.erro(job_id, f"Erro durante o processamento: {str(e)}")
        raise e

# ---------------------------------------------------------
# Endpoints da API
# ---------------------------------------------------------
@app.get("/", tags=["Status"])
def root():
    return {
        "status": "online",
        "api": "Gerador de Audiolivro e Sincronizacao de Texto",
        "versao": "2.0.0",
        "docs": "/docs",
        "health": "/health",
        "vozes": "/vozes"
    }

@app.get("/health", tags=["Status"])
def health_check():
    return {"status": "ok", "timestamp": time.time()}

@app.get("/vozes", tags=["Configuração"])
def listar_vozes():
    """Retorna uma lista de vozes recomendadas em Português."""
    return {
        "vozes_recomendadas": [
            {"id": "pt-BR-AntonioNeural", "genero": "Masculino", "idioma": "pt-BR", "padrao": True},
            {"id": "pt-BR-FranciscaNeural", "genero": "Feminino", "idioma": "pt-BR"},
            {"id": "pt-BR-ThalitaNeural", "genero": "Feminino", "idioma": "pt-BR"},
            {"id": "pt-PT-DuarteNeural", "genero": "Masculino", "idioma": "pt-PT"},
            {"id": "pt-PT-RaquelNeural", "genero": "Feminino", "idioma": "pt-PT"}
        ]
    }

@app.post("/processar-livro/", tags=["Processamento"])
async def processar_livro(
    request: Request,
    background_tasks: BackgroundTasks,
    file: UploadFile = File(..., description="Arquivo .pdf ou .docx do livro/documento"),
    termo_cabecalho: Optional[str] = Form(None, description="Cabeçalho ou termo recorrente a ignorar (opcional)"),
    voz: str = Form("pt-BR-AntonioNeural", description="Voz do Edge-TTS"),
    rate: str = Form("+0%", description="Taxa de velocidade (ex: +0%, +10%, -5%)"),
    pitch: str = Form("-5Hz", description="Ajuste de tom (ex: -5Hz, +0Hz)"),
    tamanho_lote: int = Form(10000, description="Tamanho máximo de caracteres por lote"),
    retornar_mapa_inline: bool = Form(False, description="Se True, inclui a lista completa de palavras e tempos diretamente no JSON de resposta"),
    sync: bool = Form(False, description="Se True, aguarda a conclusão antes de responder (modo bloqueante). Se False (padrão), retorna job_id imediatamente para streaming via SSE.")
):
    """
    Recebe um arquivo PDF ou DOCX, extrai e limpa o texto,
    sintetiza a narração sincronizada via Edge-TTS e transmite
    logs/progresso em tempo real via SSE.
    """
    extensao = Path(file.filename or "").suffix.lower()
    if extensao not in [".pdf", ".docx"]:
        raise HTTPException(
            status_code=400,
            detail="Formato de arquivo não suportado. Envie um arquivo .pdf ou .docx."
        )

    conteudo_bytes = await file.read()
    if not conteudo_bytes:
        raise HTTPException(status_code=400, detail="O arquivo enviado está vazio.")

    # Limpeza periódica de armazenamento
    background_tasks.add_task(limpar_arquivos_antigos)

    # Identificador único do job
    job_id = uuid.uuid4().hex
    nome_original = Path(file.filename or "livro").stem
    base_url = str(request.base_url).rstrip("/")

    # Registra o job no gerenciador de streaming
    log_manager.register_job(job_id)

    if sync:
        # Modo bloqueante legado
        return await executar_processamento_livro(
            job_id=job_id,
            conteudo_bytes=conteudo_bytes,
            extensao=extensao,
            nome_original=nome_original,
            termo_cabecalho=termo_cabecalho,
            voz=voz,
            rate=rate,
            pitch=pitch,
            tamanho_lote=tamanho_lote,
            retornar_mapa_inline=retornar_mapa_inline,
            base_url=base_url
        )

    # Modo assíncrono padrão com streaming SSE
    asyncio.create_task(
        executar_processamento_livro(
            job_id=job_id,
            conteudo_bytes=conteudo_bytes,
            extensao=extensao,
            nome_original=nome_original,
            termo_cabecalho=termo_cabecalho,
            voz=voz,
            rate=rate,
            pitch=pitch,
            tamanho_lote=tamanho_lote,
            retornar_mapa_inline=retornar_mapa_inline,
            base_url=base_url
        )
    )

    return {
        "status": "iniciado",
        "job_id": job_id,
        "stream_url": f"{base_url}/jobs/{job_id}/stream",
        "status_url": f"{base_url}/status/{job_id}",
        "download_audio_url": f"{base_url}/download/{job_id}/audio",
        "download_mapa_url": f"{base_url}/download/{job_id}/mapa"
    }

@app.get("/jobs/{job_id}/stream", tags=["Streaming"])
async def stream_logs_job(job_id: str):
    """
    Túnel SSE (Server-Sent Events) para transmissão em tempo real
    dos logs e porcentagem de progresso de um job específico.
    """
    return StreamingResponse(
        log_manager.stream_events(job_id),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no"
        }
    )

@app.get("/download/{job_id}/audio", tags=["Download"])
def baixar_audio(job_id: str):
    """Retorna o arquivo de áudio MP3 compilado para reprodução ou download."""
    job_dir = STORAGE_DIR / job_id
    if not job_dir.exists():
        raise HTTPException(status_code=404, detail="Job não encontrado ou expirado.")

    arquivos_mp3 = list(job_dir.glob("*_Audiolivro.mp3"))
    if not arquivos_mp3:
        raise HTTPException(status_code=404, detail="Arquivo de áudio não encontrado.")

    arquivo_mp3 = arquivos_mp3[0]
    return FileResponse(
        path=arquivo_mp3,
        media_type="audio/mpeg",
        filename=arquivo_mp3.name
    )

@app.get("/download/{job_id}/mapa", tags=["Download"])
def baixar_mapa(job_id: str):
    """Retorna o arquivo JSON com o mapa de sincronização de tempos e palavras."""
    job_dir = STORAGE_DIR / job_id
    if not job_dir.exists():
        raise HTTPException(status_code=404, detail="Job não encontrado ou expirado.")

    arquivos_mapa = list(job_dir.glob("*_MapaDoTempo.json"))
    if not arquivos_mapa:
        raise HTTPException(status_code=404, detail="Arquivo de mapa JSON não encontrado.")

    arquivo_mapa = arquivos_mapa[0]
    return FileResponse(
        path=arquivo_mapa,
        media_type="application/json",
        filename=arquivo_mapa.name
    )

@app.get("/status/{job_id}", tags=["Status"])
def status_job(job_id: str):
    """Consulta os metadados ou status atual de um processamento específico."""
    job_dir = STORAGE_DIR / job_id
    arquivo_info = job_dir / "info.json"

    if arquivo_info.exists():
        with open(arquivo_info, "r", encoding="utf-8") as f:
            dados = json.load(f)
            return {"status": "concluido", **dados}

    if job_id in log_manager.active_jobs or job_id in log_manager.job_history:
        return {
            "status": log_manager.job_status.get(job_id, "processando"),
            "job_id": job_id,
            "logs_recentes": log_manager.job_history.get(job_id, [])[-15:]
        }

    raise HTTPException(status_code=404, detail="Job não encontrado ou expirado.")

if __name__ == "__main__":
    import uvicorn
    # Permite rodar diretamente via 'python app.py' para testes locais
    uvicorn.run("app:app", host="0.0.0.0", port=7860, reload=True)
